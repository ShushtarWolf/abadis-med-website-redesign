#!/usr/bin/env python3
"""Build site/assets/data/iran-hospitals-centers.json from the hospitals workbook.

The hospital sheet has no postal-code column. The surgery sheet may still
carry an empty «کد پستی» column. Columns are read by header name, and postal
codes are never written into the JSON.

Default input: tools/redesign/data/hospitals-excel-2026-10-27.xlsx
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_XLSX = ROOT / 'tools' / 'redesign' / 'data' / 'hospitals-excel-2026-10-27.xlsx'
OUT = ROOT / 'site' / 'assets' / 'data' / 'iran-hospitals-centers.json'

# Header names after ZWNJ/ZWJ are stripped. Postal code is intentionally absent.
NAME = 'نام مرکز'
PROVINCE = 'استان'
CITY = 'شهر یا شهرستان'
CITY_ALT = 'شهر'
KIND = 'نوع ثبت در منبع'
LON = 'طول جغرافیایی'
LAT = 'عرض جغرافیایی'
OWNER = 'نوع مالکیت'
OWNER_ALT = 'مالکیت'
OSM_ID = 'شناسه منبع'
OSM_URL = 'لینک منبع'
POSTAL = 'کد پستی'
UNKNOWN = 'نامشخص'


def norm_header(value) -> str:
    if value is None:
        return ''
    return str(value).replace('\u200c', '').replace('\u200d', '').strip()


def cell(value):
    if value is None:
        return ''
    text = str(value).strip()
    if text.lower() == 'none':
        return ''
    return text


def sheet_type(title: str) -> str | None:
    name = norm_header(title)
    if 'جراحی' in name and 'بیمارستان' not in name:
        return 'surgery'
    if 'بیمارستان' in name:
        return 'hospital'
    return None


def index_headers(header) -> dict:
    found = {}
    for i, raw in enumerate(header):
        key = norm_header(raw)
        if key and key not in found:
            found[key] = i
    return found


def need(headers: dict, *names: str) -> int:
    for name in names:
        if name in headers:
            return headers[name]
    raise SystemExit(f'missing column {names[0]!r}; headers={list(headers)}')


def load_rows(path: Path) -> list[dict]:
    try:
        import openpyxl
    except ImportError:
        raise SystemExit('openpyxl is required: pip install openpyxl')
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    centers = []
    seen = set()
    for title in wb.sheetnames:
        kind = sheet_type(title)
        ws = wb[title]
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            continue
        headers = index_headers(header)
        headers.pop(POSTAL, None)
        if NAME not in headers or PROVINCE not in headers:
            continue
        if kind is None:
            kind = 'hospital'
        name_i = headers[NAME]
        prov_i = headers[PROVINCE]
        city_i = need(headers, CITY, CITY_ALT)
        lon_i = need(headers, LON)
        lat_i = need(headers, LAT)
        own_i = headers.get(OWNER, headers.get(OWNER_ALT))
        id_i = headers.get(OSM_ID)
        url_i = headers.get(OSM_URL)
        # Postal code is ignored even when the sheet still has the column.
        for raw in rows:
            if raw is None:
                continue
            name = cell(raw[name_i] if name_i < len(raw) else '')
            province = cell(raw[prov_i] if prov_i < len(raw) else '')
            city = cell(raw[city_i] if city_i < len(raw) else '')
            if not name or not province:
                continue
            try:
                lon = round(float(raw[lon_i]), 6)
                lat = round(float(raw[lat_i]), 6)
            except (TypeError, ValueError):
                continue
            ownership = cell(raw[own_i]) if own_i is not None and own_i < len(raw) else ''
            osm_id = cell(raw[id_i]) if id_i is not None and id_i < len(raw) else ''
            osm = cell(raw[url_i]) if url_i is not None and url_i < len(raw) else ''
            if not osm and osm_id:
                osm = 'https://www.openstreetmap.org/' + osm_id.lstrip('/')
            key = osm_id or osm or (name, province, city, kind, lon, lat)
            if key in seen:
                continue
            seen.add(key)
            item = {
                'name': name,
                'province': province,
                'city': city,
                'type': kind,
                'lon': lon,
                'lat': lat,
                'city_fa': city,
                'ownership': ownership,
            }
            if osm:
                item['osm'] = osm
            centers.append(item)
    wb.close()
    return centers


def build(path: Path) -> dict:
    centers = load_rows(path)
    hospitals = sum(1 for c in centers if c['type'] == 'hospital')
    surgery = sum(1 for c in centers if c['type'] == 'surgery')
    known = lambda t: sum(
        1 for c in centers
        if c['type'] == t and c.get('ownership') and c['ownership'] != UNKNOWN
    )
    if any('postal' in c or 'کد پستی' in c for c in centers):
        raise SystemExit('postal code leaked into a center record')
    return {
        'dataset': 'all',
        'source': path.name,
        'counts': {
            'hospital': hospitals,
            'surgery': surgery,
            'total': hospitals + surgery,
        },
        'ownership_known': {'hospital': known('hospital'), 'surgery': known('surgery')},
        'centers': centers,
    }


def main(argv: list[str]) -> None:
    src = Path(argv[1]) if len(argv) > 1 else DEFAULT_XLSX
    if not src.is_file():
        raise SystemExit(f'workbook not found: {src}')
    data = build(src)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(data, ensure_ascii=False, separators=(',', ':')) + '\n', encoding='utf-8')
    counts = data['counts']
    known = data['ownership_known']
    print(f"{OUT.relative_to(ROOT)}  hospital={counts['hospital']} surgery={counts['surgery']} total={counts['total']}")
    print(f"ownership known  hospital={known['hospital']} surgery={known['surgery']}")


if __name__ == '__main__':
    main(sys.argv)
